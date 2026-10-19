"""월간 성과 보고서: 합성 데이터로 파싱·집계·급변/채널 이동·광고 상태·렌더·병합 검증."""
import io
import os
import sys
import unittest
from datetime import date, timedelta

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from core.render import extract_report_data  # noqa: E402
from monthly.analysis import MonthlySettings, build_context, channel_groups, name_key, trend_word  # noqa: E402
from monthly.data import AdRaw, InsightData, read_ad_raw, read_insight  # noqa: E402
from monthly.render import merge_monthly, render_monthly  # noqa: E402

DAILY_HDR = ["날짜", "방문자", "조회", "장바구니", "주문", "구매전환율", "판매량", "매출(원)"]
OPT_HDR = ["옵션 ID", "옵션명", "상품명", "등록상품ID", "카테고리", "판매방식", "매출(원)", "주문", "판매량", "방문자"]


def xlsx(sheets: dict) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def month_file(y, m, days, rev_per_day, options, vis=1000, cvr="4.00%"):
    daily = [DAILY_HDR]
    for i in range(days):
        d = date(y, m, 1) + timedelta(i)
        daily.append([d.isoformat(), str(vis), "0", "0", "40", cvr, "0", str(rev_per_day)])
    daily.append(["", str(vis * days), "", "", "", "", "", str(rev_per_day * days)])  # 총계 행
    opts = [OPT_HDR] + [[oid, name, name, "1", "가구", ch, str(rev), "1", "1", "10"] for oid, name, ch, rev in options]
    return xlsx({"기간별 판매량": daily, "상품별 판매량": opts})


def ad_raw_file(rows):
    hdr = ["날짜", "캠페인", "광고집행 옵션ID", "광고비", "총 전환매출액(14일)"]
    return xlsx({"raw": [hdr] + [[d, "c", oid, cost, rev] for d, oid, cost, rev in rows]})


class MonthlyTest(unittest.TestCase):
    def setUp(self):
        # 8월(31일, 일 100만) · 9월(30일, 일 100만) · 10월 1~10일(일 100만) + 10/11 방문자 0(부분집계)
        aug = [("1", "평판등 640, 주광색", "판매자배송", 3_000_000), ("2", "평판등 640, 주광색", "로켓그로스", 0),
               ("3", "센서등", "로켓그로스", 1_000_000), ("4", "기타", "판매자배송", 27_000_000)]
        sep = [("1", "평판등 640, 주광색", "판매자배송", 1_000_000), ("2", "평판등 640, 주광색", "로켓그로스", 200_000),
               ("3", "센서등", "로켓그로스", 1_000_000), ("4", "기타", "판매자배송", 27_800_000)]
        octo = [("1", "평판등 640, 주광색", "판매자배송", 0), ("2", "평판등 640, 주광색", "로켓그로스", 1_000_000),
                ("3", "센서등", "로켓그로스", 0), ("4", "기타", "판매자배송", 9_000_000)]
        self.ins = InsightData()
        read_insight(month_file(2026, 8, 31, 1_000_000, aug), "A_8월.xlsx", self.ins)
        read_insight(month_file(2026, 9, 30, 1_000_000, sep, cvr="5.00%"), "A_9월.xlsx", self.ins)
        read_insight(month_file(2026, 10, 10, 1_000_000, octo), "A_10월.xlsx", self.ins)
        self.ins.daily[date(2026, 10, 11)] = {"vis": 0, "ord": 5, "cvr": 0, "rev": 300_000}
        # 광고 raw: 8/1~8/20, 9/10~10/10. 옵션3은 9/25 마지막 집행
        rows = []
        for d in [date(2026, 8, 1) + timedelta(i) for i in range(20)] + \
                 [date(2026, 9, 10) + timedelta(i) for i in range(31)]:
            rows.append((d, "4", 50_000, 300_000))
            if d <= date(2026, 9, 25):
                rows.append((d, "3", 10_000, 0))
        self.raw = read_ad_raw(ad_raw_file(rows), "raw.xlsx", AdRaw())
        self.s = MonthlySettings(account="테스트")

    def test_parse_skips_total_row_and_cutoff(self):
        self.assertEqual(len([d for d in self.ins.daily if d.month == 8]), 31)
        ctx = build_context(self.ins, self.raw, None, self.s)
        self.assertEqual(ctx["cutoff"], date(2026, 10, 10))  # 방문자 0인 10/11 제외
        self.assertEqual(ctx["months"][0]["rev"], 31_000_000)

    def test_same_period_and_cvr(self):
        ctx = build_context(self.ins, self.raw, None, self.s)
        self.assertTrue(ctx["last"]["partial"])
        self.assertEqual(ctx["same"]["rev"], 10_000_000)       # 9/1~10
        self.assertAlmostEqual(ctx["months"][1]["cvr"], 5.0)   # 방문자 가중평균
        rev_kpi = ctx["kpis"][0]
        self.assertEqual(rev_kpi["delta"]["value"], 0)         # 10/1~10 vs 9/1~10 동일

    def test_ad_rows_from_raw_when_no_dashboard(self):
        ctx = build_context(self.ins, self.raw, None, self.s)
        aug, sep = ctx["months"][0], ctx["months"][1]
        self.assertFalse(aug["ad_ok"])                          # 20/31일만 → 광고 지표 비움
        self.assertFalse(sep["ad_ok"])                          # 21/30일
        self.assertTrue(ctx["last"]["ad_ok"])
        self.assertEqual(ctx["last"]["cost"], 10 * 50_000)

    def test_manual_fill_makes_month_usable(self):
        manual = [(date(2026, 8, 21) + timedelta(i), 1, 1) for i in range(11)]
        ctx = build_context(self.ins, self.raw, None, MonthlySettings(account="t", manual_ads=manual))
        self.assertTrue(ctx["months"][0]["ad_ok"])

    def test_channel_shift_simultaneous_and_lagged(self):
        ctx = build_context(self.ins, self.raw, None, self.s)
        mv89 = {m["oid"]: m for m in ctx["movers"][0]["rows"]}
        self.assertIn("1", mv89)                                # -200만
        self.assertTrue(mv89["1"]["shift"].startswith("판매자배송→로켓그로스"))
        self.assertIn("시차", mv89["1"]["note"])                 # 같은 기간 +20만은 부족, 8→10 전체로 판단
        mv910 = {m["oid"]: m for m in ctx["movers"][1]["rows"]}
        self.assertIn("2", mv910)                               # 그로스 급증 (월환산)
        self.assertTrue(mv910["2"]["shift"])

    def test_ad_status(self):
        ctx = build_context(self.ins, self.raw, None, self.s)
        mv910 = {m["oid"]: m for m in ctx["movers"][1]["rows"]}
        self.assertTrue(mv910["3"]["ad_status"].startswith("중단(마지막 9/25)"))
        self.assertGreater(mv910["3"]["ad_a"], 0)
        self.assertEqual(mv910["3"]["ad_b"], 0)

    def test_status_at_data_gap_is_flagged(self):
        rows = [(date(2026, 8, 1) + timedelta(i), "9", 1000, 0) for i in range(20)] + \
               [(date(2026, 9, 10) + timedelta(i), "4", 1000, 0) for i in range(31)]
        raw = read_ad_raw(ad_raw_file(rows), "r.xlsx", AdRaw())
        from monthly.analysis import ad_context
        ad = ad_context(raw, date(2026, 10, 10), 7)
        self.assertTrue(ad["status"]("9").startswith("확인필요"))
        self.assertEqual(ad["status"]("4"), "집행중")
        self.assertEqual(ad["status"]("x"), "미집행")

    def test_option_sheet_pick_and_warning(self):
        ins = InsightData()
        ins.daily, ins.options = dict(self.ins.daily), list(self.ins.options)
        bad = month_file(2026, 8, 31, 1_000_000, [("1", "x", "판매자배송", 90_000_000)])
        wb_bad = xlsx({"상품별 판매량": openpyxl.load_workbook(io.BytesIO(bad))["상품별 판매량"].values})
        read_insight(wb_bad, "A_8월 상품별.xlsx", ins)
        ctx = build_context(ins, None, None, self.s)
        self.assertFalse(any("8월 상품별" in t for _, t in ctx["checks"]))  # 일별과 맞는 시트를 골라 경고 없음

    def test_name_key_and_groups(self):
        self.assertEqual(name_key("[전국설치] 안트 LED, JK_주광색"), name_key("안트 LED, 주광색"))
        g = channel_groups([{"1": {"name": "A", "ch": "판매자배송"}, "2": {"name": "A", "ch": "로켓그로스"},
                             "3": {"name": "B", "ch": "판매자배송"}}], [("3", "9")])
        self.assertEqual(g["1"], g["2"])
        self.assertEqual(g["3"], g["9"])

    def test_trend_word(self):
        self.assertEqual(trend_word([1, 2, 3]), "상승 중")
        self.assertEqual(trend_word([3, 2, 1]), "하락 중")
        self.assertEqual(trend_word([1, 2]), "판단 보류")

    # ---------- 날짜 없는 기간 합계 광고 파일 ----------
    def test_period_from_name(self):
        from monthly.data import period_from_name
        self.assertEqual(period_from_name("안트_9월 광고 보고서.xlsm"), (9, 1, 9, 0))
        self.assertEqual(period_from_name("안트_AI_상품전환_0901-0909.xlsx"), (9, 1, 9, 9))
        self.assertEqual(period_from_name("광고_9월16~9월22일.xlsx"), (9, 16, 9, 22))
        self.assertIsNone(period_from_name("광고보고서.xlsx"))

    def _period_file(self, rows, campaigns=("A", "B"), conv_style=False):
        if conv_style:  # 광고센터 상품 전환 보고서 헤더
            hdr = ["캠페인 이름", "광고 집행 옵션 ID", "광고비(원)", "총 전환 매출액 (14일)(원)"]
        else:           # 광고보고서 xlsm raw 시트 (날짜 칸은 있지만 비어 있음)
            hdr = ["날짜", "캠페인", "광고집행 옵션ID", "광고비", "총 전환매출액(14일)"]
        body = []
        for i, (oid, cost, rev) in enumerate(rows):
            camp = campaigns[i % len(campaigns)]
            body.append([camp, oid, cost, rev] if conv_style else [None, camp, oid, cost, rev])
        return xlsx({"raw" if not conv_style else "Sheet1": [hdr] + body})

    def test_period_file_parsed_and_used_for_totals(self):
        raw = read_ad_raw(self._period_file([("4", 2_000_000, 9_000_000), ("3", 100_000, 0)]), "A_9월 광고 보고서.xlsm", AdRaw())
        self.assertEqual(len(raw.periods), 1)
        self.assertTrue(raw.periods[0].use_total)
        dash = [{"s": date(2026, 9, 1), "e": date(2026, 9, 30), "cost": 1_000_000, "adrev": 1}]  # 틀린 대시보드
        # 대시보드는 같은 주에 캠페인별 여러 행 → 기간 합계와 안 겹치는 행은 모두 유지돼야 함
        dash += [{"s": date(2026, 10, 1), "e": date(2026, 10, 10), "cost": c, "adrev": 0} for c in (100, 200, 300)]
        ctx = build_context(self.ins, raw, dash, self.s)
        sep = ctx["months"][1]
        self.assertTrue(sep["ad_ok"])
        self.assertEqual(sep["cost"], 2_100_000)                # 광고보고서가 대시보드보다 우선
        self.assertEqual(ctx["last"]["cost"], 600)              # 10월 캠페인 3행 합계
        self.assertTrue(any("대시보드" in t and "다릅니다" in t for _, t in ctx["checks"]))

    def test_conversion_report_headers_and_single_campaign(self):
        raw = read_ad_raw(self._period_file([("3", 50_000, 0)], campaigns=("AI",), conv_style=True),
                          "AI_상품전환_0901-0909.xlsx", AdRaw())
        p = raw.periods[0]
        self.assertFalse(p.use_total)                           # 캠페인 1개 → 계정 전체 아님
        ctx = build_context(self.ins, raw, None, self.s)
        self.assertFalse(ctx["months"][1]["ad_ok"])             # 합계에는 안 씀
        from monthly.analysis import ad_context, resolve_periods
        resolve_periods(raw, self.ins.daily)
        ad = ad_context(raw, date(2026, 10, 10), 7)
        self.assertEqual(ad["spend"]("3", date(2026, 9, 1), date(2026, 9, 30)), 50_000)

    def test_period_overrides_daily_within_range_and_status(self):
        raw = AdRaw()
        read_ad_raw(ad_raw_file([(date(2026, 9, 20), "3", 10, 0), (date(2026, 10, 1), "4", 10, 0)]), "d.xlsx", raw)
        read_ad_raw(self._period_file([("3", 300_000, 0), ("4", 1, 0)]), "A_9월 광고 보고서.xlsm", raw)
        from monthly.analysis import ad_context, resolve_periods
        resolve_periods(raw, self.ins.daily)
        ad = ad_context(raw, date(2026, 10, 1), 1)
        self.assertEqual(ad["spend"]("3", date(2026, 9, 1), date(2026, 9, 30)), 300_000)  # 9/20 일별은 기간합계에 포함
        self.assertEqual(ad["coverage"](date(2026, 9, 1), date(2026, 9, 30)), (30, 30))
        self.assertEqual(ad["status"]("4"), "집행중")
        self.assertTrue(ad["status"]("3").startswith("중단(9/1~9/30 집행 후)"))

    def test_cvr_is_orders_over_views_when_views_present(self):
        rows = [DAILY_HDR] + [[f"2026-08-{d:02d}", "100", "400", "0", "20", "9.99%", "0", "1000"] for d in range(1, 11)]
        ins = read_insight(xlsx({"기간별 판매량": rows}), "B_8월.xlsx", InsightData())
        ctx = build_context(ins, None, None, MonthlySettings(account="b"))
        self.assertAlmostEqual(ctx["months"][0]["cvr"], 5.0)   # 200 ÷ 4000, 일별 칸(9.99%)이 아님

    def test_top_sheet_is_partial_and_skips_total_warning(self):
        ins = InsightData()
        ins.daily = dict(self.ins.daily)
        for m, rev in ((8, 3_000_000), (9, 5_000_000)):
            opts = [OPT_HDR, ["7", "탑상품", "탑상품", "1", "x", "로켓배송", str(rev), "1", "1", "10"]]
            read_insight(xlsx({"상품별 판매량(TOP10)": opts}), f"G_{m}월.xlsx", ins)
        self.assertTrue(all(o.partial for o in ins.options))
        ctx = build_context(ins, None, None, self.s)
        self.assertFalse(any("차이 납니다" in t for _, t in ctx["checks"]))
        self.assertTrue(any("상위 상품만" in t for _, t in ctx["checks"]))
        self.assertTrue(ctx["movers"][0]["top_only"])
        self.assertEqual(ctx["movers"][0]["rows"][0]["diff"], 2_000_000)
        # 진행 중인 달 TOP 시트는 일수 환산(31/10)만 적용
        opts = [OPT_HDR, ["7", "탑상품", "탑상품", "1", "x", "로켓배송", "1000000", "1", "1", "10"]]
        read_insight(xlsx({"상품별 판매량(TOP10)": opts}), "G_10월.xlsx", ins)
        ctx = build_context(ins, None, None, self.s)
        mv = ctx["movers"][1]["rows"][0]
        self.assertAlmostEqual(mv["b"], 3_100_000)

    def test_render_and_merge(self):
        ctx = build_context(self.ins, self.raw, None, self.s)
        html = render_monthly(ctx)
        for part in ("월별 핵심 지표", "① 매출 구성", "② 최근 흐름", "③ 매출 급변 상품", 'id="report-data"'):
            self.assertIn(part, html)
        data = extract_report_data(html)
        self.assertEqual(data["kind"], "monthly")
        other = html.replace("테스트", "다른계정")
        merged, skipped = merge_monthly([{"label": "A", "html": html}, {"label": "B", "html": other}], "통합")
        self.assertEqual(skipped, [])
        self.assertIn("계정 통합 보고서", merged)
        self.assertNotIn("</script><", merged.split('id="embedded-reports"')[1][:200])


if __name__ == "__main__":
    unittest.main()
