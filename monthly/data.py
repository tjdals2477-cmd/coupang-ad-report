"""월간 성과 보고서 입력 파일 읽기.

- 쿠팡 Wing 비즈니스 인사이트: '기간별 판매량'(일별) + '상품별 판매량'(옵션별) 시트. 상품별만 따로 받은 파일도 가능.
- 쿠팡 광고 raw: 일자 × 광고집행 옵션ID 행 (광고센터 보고서 또는 광고보고서 xlsm의 'raw' 시트)
- 광고 주차 대시보드(선택): 시작일·종료일·광고비용·광고매출 행 (시트 = 계정)
시트·컬럼은 이름으로 찾으므로 순서가 바뀌어도 된다.
"""
from __future__ import annotations

import io
import re
import warnings
from dataclasses import dataclass, field
from datetime import date, datetime

import openpyxl

warnings.filterwarnings("ignore", module="openpyxl")


class MonthlyFileError(ValueError):
    pass


def num(v) -> float:
    if v in (None, "", "-"):
        return 0.0
    try:
        return float(str(v).replace(",", "").replace("%", "").strip())
    except ValueError:
        return 0.0


def to_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        m = re.match(r"\s*(\d{4})[-./](\d{1,2})[-./](\d{1,2})", v)
        if m:
            return date(int(m[1]), int(m[2]), int(m[3]))
    return None


def _month_from_name(name: str) -> int | None:
    m = re.search(r"(\d{1,2})\s*월", name)
    return int(m[1]) if m and 1 <= int(m[1]) <= 12 else None


def _open(data: bytes):
    try:
        return openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # 손상·비엑셀 파일
        raise MonthlyFileError(f"엑셀 파일을 열 수 없습니다: {e}") from e


def _header(row) -> list[str]:
    return [str(c).strip() if c is not None else "" for c in row]


# ---------------- 비즈니스 인사이트 ----------------

DAILY_COLS = {"날짜", "방문자", "주문", "구매전환율", "매출(원)"}
OPTION_COLS = {"옵션 ID", "옵션명", "판매방식", "매출(원)", "방문자"}


@dataclass
class OptionSheet:
    source: str            # "파일명 · 시트명"
    month: int | None
    rows: dict             # {옵션ID: {name, ch, rev, vis, ord}}
    partial: bool = False  # 시트명에 'TOP'이 있으면 상위 상품만 담긴 시트 (합계 대조·전체 비교 불가)

    @property
    def total(self) -> float:
        return sum(r["rev"] for r in self.rows.values())


@dataclass
class InsightData:
    daily: dict = field(default_factory=dict)     # {date: {vis, view, ord, cvr, rev}}
    options: list = field(default_factory=list)   # [OptionSheet]
    notes: list = field(default_factory=list)


def read_insight(data: bytes, name: str, into: InsightData | None = None) -> InsightData:
    out = into or InsightData()
    wb = _open(data)
    found = False
    daily_month = None
    pending_opts = []
    for ws in wb.worksheets:
        it = ws.iter_rows(values_only=True)
        try:
            hdr = _header(next(it))
        except StopIteration:
            continue
        if DAILY_COLS <= set(hdr):
            ix = {h: hdr.index(h) for h in DAILY_COLS}
            i_view = hdr.index("조회") if "조회" in hdr else None
            months = []
            for r in it:
                d = to_date(r[ix["날짜"]])
                if not d:
                    continue  # 하단 총계 행
                out.daily[d] = {"vis": num(r[ix["방문자"]]), "ord": num(r[ix["주문"]]),
                                "view": num(r[i_view]) if i_view is not None else 0.0,
                                "cvr": num(r[ix["구매전환율"]]), "rev": num(r[ix["매출(원)"]])}
                months.append(d.month)
            if months:
                daily_month = max(set(months), key=months.count)
                found = True
        elif OPTION_COLS <= set(hdr):
            ix = {h: hdr.index(h) for h in OPTION_COLS}
            i_ord = hdr.index("주문") if "주문" in hdr else None
            rows = {}
            for r in it:
                oid = r[ix["옵션 ID"]]
                if oid in (None, ""):
                    continue
                oid = str(oid).split(".")[0].strip()
                rows[oid] = {"name": str(r[ix["옵션명"]] or ""), "ch": str(r[ix["판매방식"]] or ""),
                             "rev": num(r[ix["매출(원)"]]), "vis": num(r[ix["방문자"]]),
                             "ord": num(r[i_ord]) if i_ord is not None else 0.0}
            pending_opts.append((ws.title, rows))
            found = True
    for title, rows in pending_opts:
        month = daily_month or _month_from_name(name)
        if month is None:
            out.notes.append(f"'{name}' 상품별 시트의 월을 알 수 없어 제외했습니다. 파일명에 'n월'을 넣어주세요.")
            continue
        out.options.append(OptionSheet(source=f"{name} · {title}", month=month, rows=rows,
                                       partial="TOP" in title.upper()))
    if not found:
        raise MonthlyFileError(f"'{name}'에서 '기간별 판매량'(날짜·방문자·매출) 또는 '상품별 판매량'(옵션 ID·판매방식·매출) 시트를 찾지 못했습니다.")
    return out


# ---------------- 광고 raw (옵션 단위) ----------------
# 광고보고서 xlsm 'raw' 시트, 광고센터 상품 전환 보고서 등 헤더 표기가 조금씩 달라 공백·(원)을 무시하고 찾는다.
AD_ALIASES = {
    "date": ("날짜", "일자"),
    "opt": ("광고집행옵션ID",),
    "cost": ("광고비", "광고비(원)"),
    "rev": ("총전환매출액(14일)", "총전환매출액(14일)(원)"),
    "camp": ("캠페인", "캠페인이름", "캠페인명"),
}


def _norm(h: str) -> str:
    return re.sub(r"\s+", "", h)


def _find(hdr: list[str], key: str) -> int | None:
    names = {_norm(a) for a in AD_ALIASES[key]}
    return next((i for i, h in enumerate(hdr) if _norm(h) in names), None)


def period_from_name(name: str) -> tuple | None:
    """파일명에서 기간 (시작월, 시작일, 종료월, 종료일). 일(day) 0은 '월말'."""
    m = re.search(r"(?<!\d)(\d{2})(\d{2})\s*[-~]\s*(\d{2})(\d{2})(?!\d)", name)       # 0901-0909
    if m:
        return int(m[1]), int(m[2]), int(m[3]), int(m[4])
    m = re.search(r"(\d{1,2})\s*월\s*(\d{1,2})\s*일?\s*[-~]\s*(?:(\d{1,2})\s*월\s*)?(\d{1,2})\s*일", name)  # 9월16~9월22일
    if m:
        return int(m[1]), int(m[2]), int(m[3] or m[1]), int(m[4])
    mo = _month_from_name(name)
    return (mo, 1, mo, 0) if mo else None


@dataclass
class AdPeriod:
    """날짜 없이 기간 합계로 내려받은 광고 파일."""
    source: str
    md: tuple | None                               # period_from_name 결과
    opt: dict = field(default_factory=dict)        # {옵션ID: [광고비, 14일 전환매출]}
    campaigns: set = field(default_factory=set)
    s: date | None = None                          # 분석 시 확정 (연도는 판매 데이터에서)
    e: date | None = None
    use_total: bool = True                         # 계정 전체 광고 합계로도 사용할지
    clamped: bool = False                          # 진행 중인 달이라 종료일을 집계 마지막 날로 맞춤

    @property
    def cost(self) -> float:
        return sum(v[0] for v in self.opt.values())

    @property
    def rev(self) -> float:
        return sum(v[1] for v in self.opt.values())


@dataclass
class AdRaw:
    by_day: dict = field(default_factory=dict)     # {date: [(옵션ID, 광고비, 14일 전환매출)]}
    periods: list = field(default_factory=list)    # [AdPeriod]


def read_ad_raw(data: bytes, name: str, into: AdRaw | None = None) -> AdRaw:
    """일별 파일: 같은 날짜는 나중에 읽은 파일이 덮어씀. 날짜가 없는 파일: 기간 합계(AdPeriod)로 보관."""
    out = into or AdRaw()
    wb = _open(data)
    sheets = [wb["raw"]] if "raw" in wb.sheetnames else wb.worksheets
    for ws in sheets:
        it = ws.iter_rows(values_only=True)
        try:
            hdr = _header(next(it))
        except StopIteration:
            continue
        i_o, i_c = _find(hdr, "opt"), _find(hdr, "cost")
        if i_o is None or i_c is None:
            continue
        i_d, i_r, i_k = _find(hdr, "date"), _find(hdr, "rev"), _find(hdr, "camp")
        dated: dict = {}
        undated = AdPeriod(source=name, md=period_from_name(name))
        for r in it:
            oid = str(r[i_o]).split(".")[0].strip() if r[i_o] not in (None, "") else ""
            cost, rev = num(r[i_c]), (num(r[i_r]) if i_r is not None else 0.0)
            d = to_date(r[i_d]) if i_d is not None else None
            if d:
                dated.setdefault(d, []).append((oid, cost, rev))
            else:
                v = undated.opt.setdefault(oid, [0.0, 0.0])
                v[0] += cost; v[1] += rev
            if i_k is not None and r[i_k] not in (None, "", "-"):
                undated.campaigns.add(str(r[i_k]))
        if dated:
            out.by_day.update(dated)
        elif undated.opt:
            undated.use_total = len(undated.campaigns) >= 2   # 캠페인 1개짜리 파일은 계정 전체가 아님
            out.periods.append(undated)
        return out
    raise MonthlyFileError(f"'{name}'에서 광고 데이터(광고집행 옵션ID·광고비 컬럼)를 찾지 못했습니다.")


# ---------------- 광고 주차 대시보드 ----------------

def dashboard_sheets(data: bytes) -> list[str]:
    return _open(data).sheetnames


def read_ad_dashboard(data: bytes, sheet: str) -> list[dict]:
    """[{s, e, cost, adrev}] — '시작일','종료일','광고비용','광고매출' 헤더가 있는 표."""
    wb = _open(data)
    if sheet not in wb.sheetnames:
        raise MonthlyFileError(f"대시보드에 '{sheet}' 시트가 없습니다.")
    rows = list(wb[sheet].iter_rows(values_only=True))
    for i, r in enumerate(rows[:15]):
        hdr = _header(r)
        if {"시작일", "종료일", "광고비용", "광고매출"} <= set(hdr):
            ix = {h: hdr.index(h) for h in ("시작일", "종료일", "광고비용", "광고매출")}
            out = []
            for r2 in rows[i + 1:]:
                s, e = to_date(r2[ix["시작일"]]), to_date(r2[ix["종료일"]])
                if s and e:
                    out.append({"s": s, "e": e, "cost": num(r2[ix["광고비용"]]), "adrev": num(r2[ix["광고매출"]])})
            return out
    raise MonthlyFileError(f"'{sheet}' 시트에서 시작일·종료일·광고비용·광고매출 헤더를 찾지 못했습니다.")
